"""Heuristic resume text extraction and parsing.

Deliberately conservative: nothing parsed here is treated as fact. Every result is stored as
'unverified' and the user confirms or corrects it. Fields that cannot be found stay empty.
(AI-assisted extraction arrives with the provider layer in a later stage.)
"""
import io
import re
from dataclasses import dataclass, field

_HEADINGS = {
    "summary": r"(professional\s+)?(summary|profile|objective|about\s+me)",
    "experience": r"(work\s+|professional\s+|relevant\s+)?(experience|employment(\s+history)?|work\s+history)",
    "education": r"education(al)?(\s+background)?|academic\s+background",
    "skills": r"(technical\s+|key\s+|core\s+)?(skills|competencies)(\s*(&|and)\s*tools)?|technologies",
    "certifications": r"(certifications?|licen[sc]es?)(\s*(&|and)\s*(certifications?|licen[sc]es?))?",
    "projects": r"(personal\s+|selected\s+|key\s+)?projects",
}
_HEAD_RES = {k: re.compile(rf"^\s*(?:{v})\s*:?\s*$", re.I) for k, v in _HEADINGS.items()}

_MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
_DATE = rf"(?:{_MONTH}\s+)?(?:19|20)\d{{2}}"
_RANGE = re.compile(rf"({_DATE})\s*(?:-|–|—|to)\s*({_DATE}|present|current|now)", re.I)
_EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_PHONE = re.compile(r"(?<!\d)(\+?\d[\d\s().-]{7,}\d)(?!\d)")
_BULLET = re.compile(r"^[\s•*\-–·▪●○◦]+")
_DEGREE = re.compile(r"\b(bachelor|master|b\.?s\.?c?|m\.?s\.?c?|ba|ma|mba|ph\.?d|diploma|associate|degree)\b", re.I)

MAX_ITEMS = 30


def extract_text(data: bytes, ext: str) -> str:
    ext = ext.lower()
    if ext == ".txt":
        return data.decode("utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages)
    if ext == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    raise ValueError(f"unsupported extension {ext}")


def split_sections(text: str) -> tuple[list[str], dict[str, list[str]]]:
    header: list[str] = []
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for raw in text.splitlines():
        line = raw.strip()
        found = None
        if line and len(line) <= 40:
            for key, rx in _HEAD_RES.items():
                if rx.match(line):
                    found = key
                    break
        if found:
            current = found
            sections.setdefault(current, [])
        elif current is None:
            header.append(line)
        else:
            sections[current].append(line)
    return header, sections


def clean(line: str) -> str:
    return _BULLET.sub("", line).strip()


def parse_skills(lines: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for line in lines:
        line = clean(line)
        if ":" in line:
            line = line.split(":", 1)[1]
        for tok in re.split(r"[,;|•·●▪]|\s{2,}", line):
            tok = tok.strip(" .\t")
            if not tok or len(tok) > 40 or len(tok.split()) > 4:
                continue
            key = tok.lower()
            if key not in seen:
                seen.add(key)
                out.append(tok)
    return out[: MAX_ITEMS * 4]


def split_blocks(lines: list[str]) -> list[list[str]]:
    """Group lines into entries: blank lines separate blocks; a second date range also starts a new block."""
    blocks: list[list[str]] = []
    cur: list[str] = []
    cur_has_date = False
    for line in lines:
        if not line:
            if cur:
                blocks.append(cur)
                cur, cur_has_date = [], False
            continue
        has_date = bool(_RANGE.search(line))
        if has_date and cur_has_date:
            blocks.append(cur)
            cur, cur_has_date = [], False
        cur.append(line)
        cur_has_date = cur_has_date or has_date
    if cur:
        blocks.append(cur)
    return blocks


def parse_employment(block: list[str]) -> dict[str, str]:
    start = end = ""
    header = ""
    rest: list[str] = []
    for line in block:
        m = _RANGE.search(line)
        if m and not start:
            start, end = m.group(1), m.group(2)
            line = _RANGE.sub("", line).strip(" ,|-–—()")
        if not header and line:
            header = line
        elif line:
            rest.append(clean(line))
    title, company = header, ""
    for sep in (" at ", " @ ", " | ", " - ", " – ", " — ", ", "):
        if sep in header:
            title, company = (s.strip() for s in header.split(sep, 1))
            break
    return {
        "title": title[:200],
        "company": company[:200],
        "start": start,
        "end": end,
        "description": "\n".join(rest)[:3000],
    }


def parse_education(block: list[str]) -> dict[str, str]:
    text = " ".join(block)
    m = _RANGE.search(text)
    years = f"{m.group(1)}-{m.group(2)}" if m else (re.findall(r"(?:19|20)\d{2}", text) or [""])[-1]
    lines = [clean(_RANGE.sub("", l)).strip(" ,|-–—()") for l in block if clean(l)]
    lines = [l for l in lines if l]
    school = lines[0] if lines else ""
    degree = lines[1] if len(lines) > 1 else ""
    if _DEGREE.search(school) and not _DEGREE.search(degree):
        school, degree = degree, school
    return {"school": school[:200], "degree": degree[:200], "years": years}


def parse_certifications(lines: list[str]) -> list[dict[str, str]]:
    out = []
    for line in lines:
        line = clean(line)
        if not line:
            continue
        years = re.findall(r"(?:19|20)\d{2}", line)
        name, issuer = line, ""
        for sep in (" - ", " | ", " – ", ", "):
            if sep in line:
                name, issuer = (s.strip() for s in line.split(sep, 1))
                break
        out.append({"name": name[:200], "issuer": re.sub(r"(?:19|20)\d{2}", "", issuer).strip(" ,-()")[:200], "year": years[-1] if years else ""})
    return out[:MAX_ITEMS]


def parse_projects(block: list[str]) -> dict[str, str]:
    lines = [clean(l) for l in block if clean(l)]
    return {"name": (lines[0] if lines else "")[:200], "summary": "\n".join(lines[1:])[:2000]}


@dataclass
class ParsedResume:
    name: str = ""
    email: str = ""
    phone: str = ""
    summary: str = ""
    employment: list[dict[str, str]] = field(default_factory=list)
    skills: list[str] = field(default_factory=list)
    education: list[dict[str, str]] = field(default_factory=list)
    certifications: list[dict[str, str]] = field(default_factory=list)
    projects: list[dict[str, str]] = field(default_factory=list)
    sections_found: list[str] = field(default_factory=list)


def parse_resume(text: str) -> ParsedResume:
    header, sections = split_sections(text)
    r = ParsedResume(sections_found=sorted(sections))
    if m := _EMAIL.search(text):
        r.email = m.group(0)
    if m := _PHONE.search("\n".join(header) or text[:600]):
        r.phone = m.group(1).strip()
    for line in header:
        words = line.split()
        if line and 1 < len(words) <= 5 and not re.search(r"[\d@]", line):
            r.name = line[:200]
            break
    r.summary = " ".join(sections.get("summary", [])).strip()[:1500]
    r.skills = parse_skills(sections.get("skills", []))
    r.employment = [parse_employment(b) for b in split_blocks(sections.get("experience", []))][:MAX_ITEMS]
    r.education = [parse_education(b) for b in split_blocks(sections.get("education", []))][:MAX_ITEMS]
    r.certifications = parse_certifications(sections.get("certifications", []))
    r.projects = [parse_projects(b) for b in split_blocks(sections.get("projects", []))][:MAX_ITEMS]
    return r
